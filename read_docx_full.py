import docx

def extract_text(doc_path):
    doc = docx.Document(doc_path)
    full_text = []
    
    # Extract paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text)
            
    # Extract tables
    for table in doc.tables:
        full_text.append("-" * 40)
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                row_data.append(cell.text.replace('\n', ' ').strip())
            full_text.append(" | ".join(row_data))
        full_text.append("-" * 40)
            
    with open(r'd:\GitHub\yt-automation\script-example-full.txt', 'w', encoding='utf-8') as f:
        f.write('\n'.join(full_text))
        
    print(f"Extraction complete. Lines: {len(full_text)}")

extract_text(r'd:\GitHub\yt-automation\script-example.docx')
