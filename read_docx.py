from docx import Document
doc = Document(r'd:\GitHub\yt-automation\script-example.docx')
with open(r'd:\GitHub\yt-automation\script-example.txt', 'w', encoding='utf-8') as f:
    for p in doc.paragraphs:
        f.write(p.text + '\n')
print(f"Done: {len(doc.paragraphs)} paragraphs")
